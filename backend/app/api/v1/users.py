import csv
import io

from fastapi import APIRouter, HTTPException, UploadFile, File, status
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pwdlib import PasswordHash
from sqlalchemy import func, select

from app.core.dependencies import CurrentUser, DbSession
from app.models.user import User
from app.schemas.user import UserCreate, UserUpdate, UserAdminResponse
from app.core.security import hash_password


router = APIRouter(prefix="/users", tags=["User Management"])
password_hasher = PasswordHash.recommended()

VALID_ROLES = {"USER", "IT_SUPPORT", "IT_MANAGER", "ADMIN"}


def require_admin(user: CurrentUser):
    if user.role != "ADMIN":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Admin permission required",
        )


def prevent_last_admin_removal(db: DbSession, target: User):
    if target.role != "ADMIN":
        return

    admin_count = db.scalar(
        select(func.count(User.id)).where(
            User.role == "ADMIN",
            User.is_active.is_(True),
        )
    ) or 0

    if target.is_active and admin_count <= 1:
        raise HTTPException(
            status_code=400,
            detail="Cannot disable or demote the last active admin",
        )


@router.get("", response_model=list[UserAdminResponse])
def list_users(db: DbSession, user: CurrentUser):
    require_admin(user)

    return db.scalars(
        select(User).order_by(User.id.desc())
    ).all()


@router.post(
    "",
    response_model=UserAdminResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_user(data: UserCreate, db: DbSession, user: CurrentUser):
    require_admin(user)

    existing = db.scalar(
        select(User).where(User.username == data.username)
    )
    if existing:
        raise HTTPException(
            status_code=409,
            detail="Username already exists",
        )

    new_user = User(
        username=data.username,
        password_hash=password_hasher.hash(data.password),
        full_name=data.full_name,
        email=str(data.email) if data.email else None,
        department=data.department.strip() if data.department else None,
        job_position=data.job_position.strip() if data.job_position else None,
        role=data.role,
        is_active=True,

    )

    db.add(new_user)
    db.commit()
    db.refresh(new_user)

    return new_user


@router.patch("/{user_id}", response_model=UserAdminResponse)
def update_user(
    user_id: int,
    data: UserUpdate,
    db: DbSession,
    user: CurrentUser,
):
    require_admin(user)

    target = db.get(User, user_id)
    if target is None:
        raise HTTPException(status_code=404, detail="User not found")

    changes = data.model_dump(exclude_unset=True)

    if user.id == target.id:
        if changes.get("is_active") is False:
            raise HTTPException(
                status_code=400,
                detail="You cannot disable your own account",
            )
        if changes.get("role") not in (None, "ADMIN"):
            raise HTTPException(
                status_code=400,
                detail="You cannot remove your own admin role",
            )

    new_role = changes.get("role", target.role)
    new_active = changes.get("is_active", target.is_active)

    if target.role == "ADMIN" and target.is_active:
        if new_role != "ADMIN" or new_active is False:
            prevent_last_admin_removal(db, target)

    if "username" in changes:
        raise HTTPException(
            status_code=400,
            detail="Username cannot be changed",
        )

    if "password" in changes:
        new_password = changes.pop("password")
        if new_password:
            target.password_hash = password_hasher.hash(new_password)
            target.must_change_password = True

    if "email" in changes:
        changes["email"] = (
            str(changes["email"]) if changes["email"] else None
        )

    for field, value in changes.items():
        setattr(target, field, value)

    db.commit()
    db.refresh(target)

    return target

@router.post("/import")
async def import_users(
    db: DbSession,
    user: CurrentUser,
    file: UploadFile = File(...),
):
    require_admin(user)

    filename = (file.filename or "").lower()

    if not filename.endswith((".xlsx", ".csv")):
        raise HTTPException(
            status_code=400,
            detail="รองรับเฉพาะไฟล์ .xlsx และ .csv",
        )

    content = await file.read()

    if not content:
        raise HTTPException(
            status_code=400,
            detail="ไฟล์ไม่มีข้อมูล",
        )

    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(
            status_code=413,
            detail="ไฟล์ต้องมีขนาดไม่เกิน 5 MB",
        )

    try:
        if filename.endswith(".csv"):
            text_content = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text_content))
            rows = list(reader)
        else:
            workbook = load_workbook(
                io.BytesIO(content),
                read_only=True,
                data_only=True,
            )
            worksheet = workbook.active
            values = worksheet.iter_rows(values_only=True)

            headers = next(values, None)

            if not headers:
                raise HTTPException(
                    status_code=400,
                    detail="ไม่พบหัวคอลัมน์ในไฟล์",
                )

            headers = [
                str(value).strip().lower() if value is not None else ""
                for value in headers
            ]

            rows = [
                dict(zip(headers, row))
                for row in values
            ]
            workbook.close()

    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="ไม่สามารถอ่านไฟล์ได้ กรุณาตรวจสอบรูปแบบไฟล์",
        )

    required = {
        "username",
        "password",
        "full_name",
        "role",
    }

    if not rows:
        raise HTTPException(
            status_code=400,
            detail="ไม่พบรายการผู้ใช้งานในไฟล์",
        )

    # ตรวจสอบชื่อคอลัมน์
    headers = {str(key).strip().lower() for key in rows[0].keys()}

    missing = required - headers

    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"ขาดคอลัมน์: {', '.join(sorted(missing))}",
        )

    valid_roles = {"USER", "IT_SUPPORT", "IT_MANAGER", "ADMIN"}

    seen_usernames = set()
    created = 0
    skipped = 0
    failed = 0
    results = []

    for row_number, row in enumerate(rows, start=2):
        row = {
            str(key).strip().lower(): value
            for key, value in row.items()
            if key is not None
        }

        username = str(row.get("username") or "").strip()
        password = str(row.get("password") or "")
        full_name = str(row.get("full_name") or "").strip()
        email = str(row.get("email") or "").strip() or None
        department = str(row.get("department") or "").strip() or None
        job_position = str(row.get("job_position") or "").strip() or None
        role = str(row.get("role") or "USER").strip().upper()

        if (
            not username
            or len(username) < 3
            or len(username) > 100
            or not full_name
            or len(full_name) > 200
            or len(password) < 8
            or len(password) > 128
            or role not in valid_roles
            or (email is not None and ("@" not in email or len(email) > 255))
            or (department is not None and len(department) > 150)
            or (job_position is not None and len(job_position) > 150)
        ):
            failed += 1
            results.append({
                "row": row_number,
                "username": username,
                "status": "failed",
                "reason": "ข้อมูลไม่ถูกต้อง หรือไม่ผ่านเงื่อนไข",
            })
            continue

        username_key = username.casefold()

        if username_key in seen_usernames:
            skipped += 1
            results.append({
                "row": row_number,
                "username": username,
                "status": "skipped",
                "reason": "Username ซ้ำภายในไฟล์",
            })
            continue

        seen_usernames.add(username_key)

        existing = db.scalar(
            select(User).where(
                func.lower(User.username) == username_key
            )
        )

        if existing:
            skipped += 1
            results.append({
                "row": row_number,
                "username": username,
                "status": "skipped",
                "reason": "Username มีอยู่ในระบบแล้ว",
            })
            continue

        new_user = User(
            username=username,
            password_hash=hash_password(password),
            full_name=full_name,
            email=email,
            department=department,
            job_position=job_position,
            role=role,
            is_active=True,
            must_change_password=True,
        )

        try:
            db.add(new_user)
            db.commit()
            created += 1
            results.append({
                "row": row_number,
                "username": username,
                "status": "created",
                "reason": "นำเข้าสำเร็จ ต้องเปลี่ยน Password เมื่อ Login",
            })
        except Exception:
            db.rollback()
            failed += 1
            results.append({
                "row": row_number,
                "username": username,
                "status": "failed",
                "reason": "ไม่สามารถบันทึกข้อมูลได้",
            })

    return {
        "total": len(rows),
        "created": created,
        "skipped": skipped,
        "failed": failed,
        "results": results,
    }

class AdminPasswordReset(BaseModel):
    password: str = Field(min_length=8, max_length=128)


@router.post("/{user_id}/reset-password")
def reset_user_password(
    user_id: int,
    data: AdminPasswordReset,
    db: DbSession,
    user: CurrentUser,
):
    require_admin(user)

    target = db.get(User, user_id)

    if target is None:
        raise HTTPException(
            status_code=404,
            detail="User not found",
        )

    target.password_hash = hash_password(data.password)
    target.must_change_password = True

    db.commit()

    return {
        "message": "Reset password successfully",
        "must_change_password": True,
    }